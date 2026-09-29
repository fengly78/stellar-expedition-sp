"""由 V1.0 设计基线生成 V1.1（批量改写标题/版本号并清理旧页码导航）。

用法：
    python doc/build_design_v11.py <V1.0.docx 的路径>

路径从命令行传入，不在源码里硬编码——本脚本原先写死了维护者本机
某个 AI 对话导出目录里的绝对路径，既泄露本机目录结构，也让公开仓库里的
这份脚本对任何人都不可运行。

产出默认写到 <仓库根>/doc/outputs/。
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Inches, RGBColor

ROOT = Path(__file__).resolve().parent.parent

if len(sys.argv) < 2:
    sys.exit('用法: python doc/build_design_v11.py <V1.0.docx>')

SRC = Path(sys.argv[1])
if not SRC.is_file():
    sys.exit(f'找不到输入文件: {SRC}')
OUT = ROOT / 'doc/outputs/New_OGame_游戏设计文档_V1.1.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)
doc=Document(SRC)
doc.paragraphs[3].text='V1 游戏设计基线与规则深化'
doc.paragraphs[6].text='版本  V1.1'
doc.paragraphs[10].text='保留 V1.0 设计基线，新增第 26 至 41 章规则深化。候选配置 NG-A0 用于原型验证，新增条款的适用关系见深化目录。'
doc.paragraphs[13].text='原版章节导航保留。新增第 26 至 41 章见文末深化目录；本版省略静态页码，使用标题导航定位。'
for p in doc.paragraphs[14:39]:
    for t in p._p.iter(qn('w:t')):
        t.text=re.sub(r'\t\d+$','',t.text or '')
    # Original hyperlink navigation remains; remove tab plus terminal number runs.
    runs=list(p._p.iter(qn('w:r')))
    found_tab=False
    for r in runs:
        if r.find(qn('w:tab')) is not None: found_tab=True
        if found_tab: r.getparent().remove(r)

for nm in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    doc.styles[nm].font.color.rgb=RGBColor(0,0,0)

# Keep original chapter text; mark superseded statements at their actual locations.
anchors=[
 ('未开始任务取消后释放全部预留','第 33 章细化预留与退款口径'),
 ('任务状态依次为草案','第 30 章细化油箱预扣与实际消耗'),
 ('同一时刻的冲突按固定顺序','第 39 章修订同刻召回判定'),
 ('撤退在轮末判断','第 31 章以确定脱离替代概率脱离'),
 ('每类可掠夺量为','第 33 章限定托管资产并明确滚动额度'),
 ('默认离线策略','第 36 章区分经济与安全授权有效期'),
 ('战后资产分为','第 34 章细化船体与保障的材料抵扣兑现'),
]
for prefix,note in anchors:
    for p in doc.paragraphs:
        if p.text.startswith(prefix) and p.style.name=='Normal':
            r=p.add_run(' 【V1.1 修订：'+note+'。】')
            r.italic=True
            r.font.size=Pt(9)
            break

lines=(ROOT/'doc/New_OGame_V1.1_设计深化.md').read_text(encoding='utf-8').splitlines()
heads=[s[3:] for s in lines if s.startswith('## ')]
doc.add_page_break()
doc.add_paragraph('V1 1 规则深化目录','Title')
doc.add_paragraph('以下内容补充原版 25 章。标记为修订的规则优先；所有新增数值均为候选配置 NG-A0。原版静态页码已移除，使用 Word 标题导航访问。')
for h in heads: doc.add_paragraph(h)

def table(block):
    rows=[[c.strip() for c in line.strip().strip('|').split('|')] for line in block]
    rows=[r for r in rows if not all(re.fullmatch(r'[-: ]+',c) for c in r)]
    t=doc.add_table(rows=1,cols=len(rows[0]))
    t.autofit=False
    total=doc.sections[0].page_width-doc.sections[0].left_margin-doc.sections[0].right_margin
    weights={3:[.20,.41,.39],4:[.25,.22,.25,.28],5:[.25,.25,.10,.21,.19],6:[.20,.23,.08,.19,.16,.14]}[len(rows[0])]
    for col,w in zip(t.columns,weights): col.width=int(total*w)
    for i,row in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        pr=t.rows[i]._tr.get_or_add_trPr()
        ns=OxmlElement('w:cantSplit'); pr.append(ns)
        if i==0:
            repeat=OxmlElement('w:tblHeader'); pr.append(repeat)
        for c,txt,w in zip(cells,row,weights):
            c.width=int(total*w)
            c.text=txt
            tcpr=c._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge); el.set(qn('w:w'),'85'); el.set(qn('w:type'),'dxa'); margins.append(el)
            tcpr.append(margins)
            if i==0:
                sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E8EDF2'); tcpr.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2)
                p.paragraph_format.space_before=Pt(2)
                p.paragraph_format.line_spacing=1.1
                for r in p.runs:
                    r.font.size=Pt(9.5)
                    r.bold=i==0
                    r.font.color.rgb=RGBColor(0,0,0)
                    r.font.name='Arial'
                    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

i=0
while i<len(lines):
    s=lines[i]
    if not s.strip() or s.startswith('# '): i+=1; continue
    if s.startswith('|'):
        b=[]
        while i<len(lines) and lines[i].startswith('|'): b.append(lines[i]); i+=1
        table(b); continue
    if s.startswith('## '):
        p=doc.add_paragraph(s[3:],'Heading 1'); p.paragraph_format.page_break_before=True
    elif s.startswith('### '): doc.add_paragraph(s[4:],'Heading 2')
    else:
        p=doc.add_paragraph(s)
        p.paragraph_format.space_after=Pt(6)
        p.paragraph_format.line_spacing=1.2
    i+=1
doc.core_properties.title='New OGame 游戏设计文档 V1.1'
doc.core_properties.subject='银河文明 MMO 规则深化与候选配置'
doc.core_properties.comments='V1.0 retained with local revision markers and chapters 26–41 added.'
doc.save(OUT)
verify=Document(OUT)
assert len([p for p in verify.paragraphs if p.style.name=='Heading 1'])==41
assert len(verify.tables)==len(Document(SRC).tables)+10
print(json.dumps({'output':str(OUT),'chapters':41,'tables':len(verify.tables),'bytes':OUT.stat().st_size},ensure_ascii=False))
